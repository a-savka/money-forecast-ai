"""Разбор банковских выписок CSV/Excel.

Поддерживаются:
- пресет Т-Банка (CSV, cp1251, разделитель «;»);
- произвольные выписки через сопоставление колонок (автоопределение по заголовкам
  или ручной маппинг из мастера импорта).

Формат колонок пресетов нужно сверять с реальными выписками — банки их меняют.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from pydantic import BaseModel


class ColumnMapping(BaseModel):
    """Номера колонок (с 0). Сумма — либо одна колонка со знаком, либо пара «расход/приход»."""

    date: int
    description: int | None = None
    amount: int | None = None
    expense: int | None = None
    income: int | None = None
    mcc: int | None = None
    category: int | None = None
    status: int | None = None


@dataclass(frozen=True)
class Preset:
    key: str
    title: str
    bank_name: str
    # Заголовки колонок (в нижнем регистре) для каждого поля
    columns: dict[str, str]
    skip_statuses: tuple[str, ...] = ()


PRESETS: dict[str, Preset] = {
    "tbank": Preset(
        key="tbank",
        title="Т-Банк (CSV)",
        bank_name="Т-Банк",
        columns={
            "date": "дата операции",
            "amount": "сумма операции",
            "description": "описание",
            "mcc": "mcc",
            "category": "категория",
            "status": "статус",
        },
        skip_statuses=("failed",),
    ),
}

# Синонимы заголовков для автоопределения маппинга
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "date": ("дата операции", "дата", "date", "дата транзакции", "transaction date"),
    "amount": ("сумма операции", "сумма", "amount", "сумма в валюте счёта", "сумма в валюте счета"),
    "expense": ("расход", "списание", "дебет", "debit", "withdrawal"),
    "income": ("приход", "поступление", "зачисление", "кредит", "credit", "deposit"),
    "description": ("описание", "назначение платежа", "description", "назначение", "получатель"),
    "mcc": ("mcc", "mcc-код", "код mcc"),
    "category": ("категория", "category"),
    "status": ("статус", "status"),
}


@dataclass
class ParsedRow:
    line: int
    date: date
    amount: Decimal
    description: str
    mcc: str | None = None
    bank_category: str | None = None


@dataclass
class ParseResult:
    headers: list[str]
    mapping: ColumnMapping | None
    preset: str | None
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    skipped: int = 0


class StatementError(ValueError):
    pass


def decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise StatementError("Не удалось определить кодировку файла")


def read_table(content: bytes, filename: str) -> list[list[str]]:
    """Возвращает таблицу строк (первая строка — заголовки)."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — openpyxl бросает разные исключения
            raise StatementError("Не удалось прочитать Excel-файл") from exc
        sheet = wb.worksheets[0]
        table = [
            ["" if v is None else _cell_to_str(v) for v in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    else:
        text = decode(content)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    table = [r for r in table if any(c.strip() for c in r)]
    if len(table) < 2:
        raise StatementError("В файле нет строк с операциями")
    return table


def _cell_to_str(v: object) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return str(v)


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", h.strip().strip('"').lower())


def detect_mapping(headers: list[str]) -> tuple[ColumnMapping | None, str | None]:
    norm = [_norm_header(h) for h in headers]

    for preset in PRESETS.values():
        if all(col in norm for col in preset.columns.values()):
            idx = {f: norm.index(col) for f, col in preset.columns.items()}
            return ColumnMapping(**idx), preset.key

    found: dict[str, int] = {}
    for fld, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:
            if syn in norm and norm.index(syn) not in found.values():
                found[fld] = norm.index(syn)
                break
    if "date" not in found or not (
        {"amount"} <= found.keys() or {"expense", "income"} & found.keys()
    ):
        return None, None
    return ColumnMapping(**found), None


DATE_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d.%m.%y",
)


def parse_date(raw: str) -> date:
    raw = raw.strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    # ISO с временем и таймзоной
    try:
        return datetime.fromisoformat(raw).date()
    except ValueError as exc:
        raise StatementError(f"Не удалось разобрать дату «{raw}»") from exc


def parse_amount(raw: str) -> Decimal | None:
    s = raw.strip().replace(" ", "").replace(" ", "").replace("−", "-").replace("–", "-")
    s = re.sub(r"(?i)(rub|руб\.?|₽)", "", s)
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")  # 1,234.56
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise StatementError(f"Не удалось разобрать сумму «{raw}»") from exc


def parse_statement(
    content: bytes, filename: str, mapping: ColumnMapping | None = None
) -> ParseResult:
    table = read_table(content, filename)
    headers, data = table[0], table[1:]
    preset_key = None
    if mapping is None:
        mapping, preset_key = detect_mapping(headers)
    result = ParseResult(headers=headers, mapping=mapping, preset=preset_key)
    if mapping is None:
        return result

    preset = PRESETS.get(preset_key) if preset_key else None
    skip_statuses = preset.skip_statuses if preset else ("failed", "отклонена", "отменена")

    def cell(row: list[str], idx: int | None) -> str:
        return row[idx].strip() if idx is not None and idx < len(row) else ""

    for n, row in enumerate(data, start=2):
        try:
            if mapping.status is not None and cell(row, mapping.status).lower() in skip_statuses:
                result.skipped += 1
                continue
            d = parse_date(cell(row, mapping.date))
            if mapping.amount is not None:
                amount = parse_amount(cell(row, mapping.amount))
            else:
                expense = parse_amount(cell(row, mapping.expense)) or Decimal(0)
                income = parse_amount(cell(row, mapping.income)) or Decimal(0)
                amount = income - abs(expense)
            if not amount:
                result.skipped += 1
                continue
            mcc = cell(row, mapping.mcc)
            result.rows.append(
                ParsedRow(
                    line=n,
                    date=d,
                    amount=amount,
                    description=cell(row, mapping.description) or "Без описания",
                    mcc=mcc if re.fullmatch(r"\d{4}", mcc) else None,
                    bank_category=cell(row, mapping.category) or None,
                )
            )
        except StatementError as exc:
            result.errors.append(f"Строка {n}: {exc}")
    return result
